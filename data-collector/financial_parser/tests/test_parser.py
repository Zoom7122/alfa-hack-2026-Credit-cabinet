import hashlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from parser import extract, money

FIXTURES = Path(os.environ.get("FINANCIAL_FIXTURES_DIR", Path(__file__).resolve().parents[1] / "testdata"))


def private_case(name):
    return json.loads((FIXTURES / "cases.json").read_text())[name]


def request(path, **changes):
    result = dict(path=str(path), applicant_type="SOLE_PROPRIETOR", inn="123456789012", tax_regime="USN_INCOME")
    if (FIXTURES / "cases.json").is_file():
        name = "girbo-2025.pdf" if Path(path).suffix.lower()==".pdf" else "kudir-2026.xls"
        if Path(path).name != "complete.xlsx":
            result.update(private_case(name)["applicant"])
    result.update(changes)
    return result


def to_xlsx(path, insert_rows=False):
    """Preserve actual merged-cell layout, optionally shifting both title and blocks."""
    import xlrd
    import openpyxl
    old = xlrd.open_workbook(FIXTURES / "kudir-2026.xls", formatting_info=True)
    new = openpyxl.Workbook(); new.remove(new.active)
    for sheet in old.sheets():
        out = new.create_sheet(sheet.name)
        def row(r):
            return r + 1 + (2 if insert_rows and r >= 7 else 0)
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                value = sheet.cell_value(r, c)
                if value != "": out.cell(row(r), c+1, value)
        for r0,r1,c0,c1 in sheet.merged_cells:
            out.merge_cells(start_row=row(r0),end_row=row(r1-1),start_column=c0+1,end_column=c1)
    new.save(path)


class RealDocuments(unittest.TestCase):
    def setUp(self):
        if self._testMethodName in {"test_numbers", "test_complete_kudir_and_continuation_sheet"}:
            return
        pdf_tests = {"test_actual_pdf_periods_units_and_negative", "test_pdf_wrong_borrower",
                     "test_document_type_not_inferred_from_filename"}
        name = "girbo-2025.pdf" if self._testMethodName in pdf_tests else "kudir-2026.xls"
        if not (FIXTURES / name).is_file() or not (FIXTURES / "cases.json").is_file():
            self.skipTest("Private fixture missing; set FINANCIAL_FIXTURES_DIR or supply local testdata/" + name)

    def test_real_xls(self):
        path = FIXTURES / "kudir-2026.xls"
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = extract(request(path))
        data = result["extracted"]
        self.assertEqual(result["document_sha256"], before)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
        self.assertEqual(data["title"]["year"]["value"], 2026)
        self.assertEqual(data["title"]["inn"]["value"], private_case("kudir-2026.xls")["applicant"]["inn"])
        self.assertEqual(data["title"]["tax_object"]["value"], "Доходы")
        self.assertEqual(len(data["operations"]), 4)
        self.assertEqual(data["operations"][2]["income"], "-126000.00")
        self.assertEqual(data["operations"][2]["date"], "2026-02-19")
        self.assertEqual(data["operations"][2]["document_number"], "30")
        self.assertEqual(data["operations"][2]["primary_document_raw"], "19.02.2026  N 30")
        self.assertEqual(data["computed_totals"]["observed_income"], "360000.00")
        self.assertEqual(data["computed_totals"]["Q1"], "360000.00")
        self.assertIsNone(data["computed_totals"]["Q2"])
        totals = {k:v["amount"] for k,v in data["declared_totals"].items()}
        self.assertEqual(totals, {"Q1":"1755000.00","Q2":"1500000.00","Q3":"1300000.00","Q4":"1800000.00",
                                  "H1":"3255000.00","M9":"4555000.00","YEAR":"6355000.00"})
        self.assertEqual(data["section_i_reference"], {})
        self.assertFalse(data["detail_complete"])
        self.assertFalse(result["data_eligible"])
        self.assertEqual({d["code"] for d in result["diagnostics"]}, {"INCOMPLETE_DETAIL", "DETAIL_REQUIRED_FOR_SCORING"})
        self.assertTrue(all(v["status"]=="PASS" for v in result["validation"] if v["code"].startswith("DECLARED_")))

    def test_xlsx_and_inserted_rows(self):
        for inserted in (False,True):
            with self.subTest(inserted=inserted), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp)/"shifted.xlsx"
                to_xlsx(path, inserted)
                result = extract(request(path))
                data = result["extracted"]
                self.assertEqual(len(data["operations"]),4)
                self.assertEqual(data["computed_totals"]["observed_income"],"360000.00")
                self.assertEqual(data["declared_totals"]["YEAR"]["amount"],"6355000.00")
                self.assertEqual(data["title"]["name"]["source"]["cell"], "AP13" if inserted else "AP11")
                self.assertEqual(data["declared_totals"]["YEAR"]["source"]["cell"], "BC69" if inserted else "BC67")

    def test_wrong_inn_blocks_every_indicator(self):
        result = extract(request(FIXTURES/"kudir-2026.xls", inn="000000000000"))
        self.assertIn("INN_MISMATCH", [d["code"] for d in result["diagnostics"]])
        self.assertTrue(all(not i["data_eligible"] for i in result["indicators"]))

    def test_other_tax_regime_unsupported(self):
        result = extract(request(FIXTURES/"kudir-2026.xls", tax_regime="USN_INCOME_MINUS_EXPENSES"))
        self.assertIn("UNSUPPORTED", [d["code"] for d in result["diagnostics"]])

    def test_formula_is_not_evaluated(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"formula.xlsx";to_xlsx(path)
            book=openpyxl.load_workbook(path);book["стр.2_3"]["BC67"]="=SUM(BC16,BC32,BC49,BC66)";book.save(path)
            result=extract(request(path))
            self.assertIn("FORMULA_UNVERIFIED", [d["code"] for d in result["diagnostics"]])
            self.assertNotIn("YEAR",result["extracted"]["declared_totals"])

    def test_filled_reference_and_zero(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"reference.xlsx";to_xlsx(path)
            book=openpyxl.load_workbook(path)
            # Unmerge the reference value region in this synthetic variant.
            s=book["стр.2_3"]
            for merged in list(s.merged_cells.ranges):
                if merged.min_row>=71: s.unmerge_cells(str(merged))
            s["BC71"]="6 355 000,00";s["BC72"]=0;book.save(path)
            result=extract(request(path))
            ref=result["extracted"]["section_i_reference"]
            self.assertEqual(ref["010"]["amount"],"6355000.00")
            self.assertEqual(ref["020"]["amount"],"0.00")
            self.assertNotIn("030",ref)

    def test_actual_pdf_periods_units_and_negative(self):
        result=extract(request(FIXTURES/"girbo-2025.pdf"))
        data=result["extracted"]
        self.assertEqual(data["taxpayer_inn"],private_case("girbo-2025.pdf")["applicant"]["inn"])
        values={(v["code"],v["period"]):v for v in data["values"]}
        self.assertEqual(values["2110","2025"]["amount"],private_case("girbo-2025.pdf")["expected"]["2110/2025"])
        self.assertEqual(values["2110","2024"]["amount"],private_case("girbo-2025.pdf")["expected"]["2110/2024"])
        self.assertEqual(values["1600","2025"]["amount"],private_case("girbo-2025.pdf")["expected"]["1600/2025"])
        self.assertEqual(values["1600","2023"]["amount"],private_case("girbo-2025.pdf")["expected"]["1600/2023"])
        self.assertEqual(values["2120","2025"]["amount"],private_case("girbo-2025.pdf")["expected"]["2120/2025"])
        self.assertEqual(values["1500","2025"]["amount"],"0.00")
        self.assertIsNone(values["1230","2024"]["amount"])
        self.assertEqual(values["2400","2025"]["source"]["page"],4)
        self.assertEqual(len(values["2400","2025"]["source"]["bbox"]),4)
        self.assertTrue(all(not i["data_eligible"] for i in result["indicators"] if i["name"]=="net_profit"))
        self.assertTrue(all(i["data_eligible"] for i in result["indicators"] if i["name"]=="accounting_revenue"))

    def test_pdf_wrong_borrower(self):
        result=extract(request(FIXTURES/"girbo-2025.pdf",applicant_type="LEGAL_ENTITY",legal_form="OOO",inn="7707329152"))
        self.assertTrue(all(not i["data_eligible"] for i in result["indicators"]))
        self.assertIn("INN_MISMATCH",[d["code"] for d in result["diagnostics"]])

    def test_document_type_not_inferred_from_filename(self):
        result=extract(request(FIXTURES/"girbo-2025.pdf",applicant_type="SOLE_PROPRIETOR"))
        self.assertEqual(result["diagnostics"][0]["code"],"UNSUPPORTED")

    def test_numbers(self):
        self.assertEqual(money(" 1\u00a0234,56 "),"1234.56")
        self.assertEqual(money("(66 099)"),"-66099.00")
        self.assertEqual(money(0),"0.00")
        for v in (None,"","…","...","—"):
            self.assertIsNone(money(v))
        for v in ("NaN","Infinity","1,001","1.234.56"):
            with self.assertRaises(ValueError): money(v)

    def test_complete_kudir_and_continuation_sheet(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"complete.xlsx"
            book=openpyxl.Workbook();title=book.active
            title.append(["Налогоплательщик", "Тестовый ИП"])
            title.append(["ИНН (для налогоплательщика - индивидуального предпринимателя)"])
            title.append(["123456789012"])
            title.append(["на 2026 год"])
            title.append(["Объект налогообложения", "Доходы"])
            first=book.create_sheet("Доходы")
            first.append(["I. Доходы и расходы"])
            first.append(["№ п/п","Дата и номер первичного документа","Содержание операции","Доходы, учитываемые при исчислении налоговой базы"])
            first.append([1,"10.01.2026 № 1","Оплата","100,00"])
            # Section I continues on a differently named sheet without another table header.
            second=book.create_sheet("Продолжение")
            second.append(["Продолжение раздела I"])
            second.append(["Итого за I квартал",None,None,100])
            second.append([2,"10.04.2026 № 2","Оплата",200])
            second.append(["Итого за II квартал",None,None,200])
            second.append(["Итого за полугодие",None,None,300])
            second.append([3,"10.07.2026 № 3","Возврат",-50])
            second.append(["Итого за III квартал",None,None,-50])
            second.append(["Итого за 9 месяцев",None,None,250])
            second.append([4,"10.10.2026 № 4","Оплата",150])
            second.append(["Итого за IV квартал",None,None,150])
            second.append(["Итого за год",None,None,400])
            book.save(path)
            result=extract(request(path))
            self.assertTrue(result["data_eligible"],result["diagnostics"])
            self.assertEqual(result["indicators"][0]["amount"],"400.00")
            self.assertEqual([o["quarter"] for o in result["extracted"]["operations"]],[1,2,3,4])
            self.assertEqual(result["extracted"]["computed_totals"]["Q3"],"-50.00")


if __name__ == "__main__": unittest.main()
