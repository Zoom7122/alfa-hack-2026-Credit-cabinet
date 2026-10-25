"""Read-only financial extraction. No signature verdicts or scoring decisions here.

JSON stdin: {path, applicant_type, inn, tax_regime}. JSON stdout: extraction report.
All amounts are decimal strings in rubles. Sources retain original text/locations.
"""
import datetime as dt
import hashlib
import json
import re
import struct
import sys
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

VERSION = "financial-v1"
MAX_CELLS = 1_000_000
MISSING = {"", "-", "—", "–", "…", "..."}


def norm(value):
    return re.sub(r"\s+", " ", str("" if value is None else value)).strip()


def money(value):
    if value is None or norm(value) in MISSING:
        return None
    text = re.sub(r"\s", "", str(value)).replace("−", "-").replace(",", ".")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise ValueError("Unrecognized amount: " + str(value))
    if not result.is_finite() or abs(result) > Decimal("1e18") or result != result.quantize(Decimal("0.01")):
        raise ValueError("Invalid monetary precision/range")
    return format(result.quantize(Decimal("0.01")), "f")


def colname(col):
    out = ""
    while col >= 0:
        out = chr(65 + col % 26) + out
        col = col // 26 - 1
    return out


class Report:
    def __init__(self):
        self.data = {"parser_version": VERSION, "extracted": {}, "validation": [],
                     "diagnostics": [], "indicators": [], "data_eligible": False}

    def issue(self, code, message, sources=None, severity="ERROR", indicators=None, period=None):
        self.data["diagnostics"].append(dict(code=code, message=message, severity=severity,
                                             sources=sources or [], indicators=indicators or [], period=period))

    def check(self, code, status, sources=None):
        self.data["validation"].append(dict(code=code, status=status, sources=sources or []))

    def finish(self):
        codes = {d["code"] for d in self.data["diagnostics"]}
        structure_errors = {"UNSUPPORTED", "PARSE_FAILED", "TITLE_FIELD_MISSING", "AMBIGUOUS_TITLE", "HEADER_UNSUPPORTED",
                            "NO_TEXT_LAYER", "FORMS_MISSING", "FORM_TABLE_MISSING", "NO_FINANCIAL_VALUES", "UNIT_UNKNOWN"}
        completeness_errors = {"DETAIL_REQUIRED_FOR_SCORING", "OPERATION_INCOMPLETE", "TOTAL_MISSING", "QUARTER_UNKNOWN",
                               "OPERATIONS_UNRECONCILED", "UNRECOGNIZED_OPERATION", "FORMULA_UNVERIFIED", "CELL_ERROR"}
        self.check("STRUCTURE", "FAIL" if codes & structure_errors else "PASS")
        self.check("COMPLETENESS", "NOT_CHECKABLE" if codes & structure_errors else ("FAIL" if codes & completeness_errors else "PASS"))
        global_errors = [x["code"] for x in self.data["diagnostics"]
                         if x["severity"] == "ERROR" and not x["indicators"]]
        for i in self.data["indicators"]:
            reasons = global_errors + [x["code"] for x in self.data["diagnostics"]
                                      if x["severity"] == "ERROR" and i["name"] in x["indicators"]
                                      and (x["period"] is None or x["period"] == i["period"])]
            i["data_eligible"] = not reasons
            i["reasons"] = sorted(set(reasons))
        self.data["data_eligible"] = bool(self.data["indicators"]) and all(
            i["data_eligible"] for i in self.data["indicators"])
        return self.data


class Sheet:
    def __init__(self, name, rows, merges):
        self.name, self.rows, self.merges = name, rows, merges

    def source(self, row, col):
        cell = colname(col) + str(row + 1)
        result = {"sheet": self.name, "cell": cell, "raw": str(self.rows[row].get(col, ""))}
        for r0, r1, c0, c1 in self.merges:
            if r0 <= row < r1 and c0 <= col < c1:
                result["range"] = f"{colname(c0)}{r0+1}:{colname(c1-1)}{r1}"
                break
        return result


def load_sheets(path, report):
    header = Path(path).read_bytes()[:8]
    if header == bytes.fromhex("d0cf11e0a1b11ae1"):
        import xlrd
        from xlrd.compdoc import CompDoc
        raw = Path(path).read_bytes()
        # xlrd exposes cached values, not formula provenance: detect BIFF FORMULA records.
        compound = CompDoc(raw)
        stream = compound.get_named_stream("Workbook") or compound.get_named_stream("Book")
        formula = False
        offset = 0
        while offset + 4 <= len(stream):
            opcode, size = struct.unpack_from("<HH", stream, offset)
            if opcode == 0x0006:
                formula = True
            offset += size + 4
        if formula:
            report.issue("FORMULA_UNVERIFIED", "XLS contains formulas; cached results cannot be verified without recalculation")
        book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
        sheets = []
        cells = 0
        for s in book.sheets():
            cells += s.nrows * s.ncols
            if cells > MAX_CELLS:
                raise ValueError("Workbook exceeds cell limit")
            rows = []
            for r in range(s.nrows):
                row = {}
                for c in range(s.ncols):
                    cell = s.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        report.issue("CELL_ERROR", "Spreadsheet error", [{"sheet": s.name, "cell": colname(c)+str(r+1)}])
                    if cell.value != "":
                        row[c] = cell.value
                rows.append(row)
            sheets.append(Sheet(s.name, rows, s.merged_cells))
        return sheets
    if header.startswith(b"PK"):
        import openpyxl
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Workbook exceeds expansion limit")
            if any("vbaproject" in x.filename.lower() for x in z.infolist()):
                raise ValueError("Macro-enabled workbooks are unsupported")
        formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, data_only=True, keep_links=False)
        sheets = []
        cells = 0
        for s in formulas:
            cells += s.max_row * s.max_column
            if cells > MAX_CELLS:
                raise ValueError("Workbook exceeds cell limit")
            rows = []
            for row in s:
                vals = {}
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        value = cached[s.title][cell.coordinate].value
                        report.issue("FORMULA_UNVERIFIED", "Formula cached value is not independently verified",
                                     [{"sheet": s.title, "cell": cell.coordinate, "raw": cell.value}])
                    if cell.data_type == "e":
                        report.issue("CELL_ERROR", "Spreadsheet error", [{"sheet": s.title, "cell": cell.coordinate}])
                    if value is not None:
                        vals[cell.column-1] = value
                rows.append(vals)
            merges = [(x.min_row-1, x.max_row, x.min_col-1, x.max_col) for x in s.merged_cells.ranges]
            sheets.append(Sheet(s.title, rows, merges))
        formulas.close()
        cached.close()
        return sheets
    raise ValueError("Not an XLS/XLSX workbook")


def title_fields(sheets, report):
    fields = {}

    def put(key, value, source):
        if key in fields and fields[key]["value"] != value:
            report.issue("AMBIGUOUS_TITLE", "Conflicting title field: " + key, [fields[key]["source"], source])
        else:
            fields[key] = {"value": value, "source": source}

    for s in sheets:
        # Title detection by labels, independent of sheet name or fixed coordinates.
        if not any("Объект налогообложения" in norm(v) for row in s.rows for v in row.values()):
            continue
        for r, row in enumerate(s.rows):
            for c, v in row.items():
                text = norm(v)
                right = [(cc, vv) for cc, vv in sorted(row.items()) if cc > c and norm(vv)]
                if text.startswith("Налогоплательщик") and right:
                    cc, vv = right[0]
                    put("name", norm(vv), s.source(r, cc))
                if text.startswith("Объект налогообложения") and right:
                    cc, vv = right[0]
                    put("tax_object", norm(vv), s.source(r, cc))
                if re.search(r"на\s+20\s*$", text, re.I):
                    for cc, vv in right:
                        if re.fullmatch(r"\d{2}(?:\.0)?", str(vv)):
                            put("year", 2000 + int(float(vv)), s.source(r, cc))
                            fields["year"]["prefix_source"] = s.source(r, c)
                            break
                elif re.fullmatch(r"на", text, re.I) and right:
                    for cc,vv in right:
                        if re.fullmatch(r"20\d{2}(?:\.0)?",str(vv)):
                            put("year",int(float(vv)),s.source(r,cc))
                            break
                elif re.fullmatch(r"на\s+20\d{2}\s*(?:год|г\.?)?", text, re.I):
                    put("year", int(re.search(r"20\d{2}", text)[0]), s.source(r, c))
                if "ИНН" in text and ("индивидуального предпринимателя" in text or text.strip(" :")=="ИНН"):
                    for rr in range(r, min(r+5, len(s.rows))):
                        digits = [(cc, str(int(float(vv)))) for cc, vv in sorted(s.rows[rr].items())
                                  if re.fullmatch(r"\d(?:\.0)?", str(vv))]
                        if len(digits) == 12:
                            source = s.source(rr, digits[0][0])
                            source["cells"] = [colname(cc)+str(rr+1) for cc, _ in digits]
                            put("inn", "".join(vv for _, vv in digits), source)
                        for cc, vv in s.rows[rr].items():
                            if re.fullmatch(r"\d{12}", norm(vv)):
                                put("inn", norm(vv), s.source(rr, cc))
    for field in ("name", "inn", "year", "tax_object"):
        if field not in fields:
            report.issue("TITLE_FIELD_MISSING", "Missing title field: " + field)
    return fields


def total_key(text):
    if "итого" not in text:
        return None
    for roman, n in (("iv", 4), ("iii", 3), ("ii", 2), ("i", 1)):
        if re.search(r"\b"+roman+r"\s+квартал", text) or re.search(r"\b"+str(n)+r"\s+квартал", text):
            return "Q" + str(n)
    if "полугод" in text:
        return "H1"
    if "9 месяц" in text or "девять месяц" in text:
        return "M9"
    if re.search(r"за\s+год", text):
        return "YEAR"
    return None


def parse_kudir(path, report, inn, tax_regime):
    sheets = load_sheets(path, report)
    title = title_fields(sheets, report)
    operations, totals, reference = [], {}, {}
    extracted = dict(kind="KUDIR", title=title, operations=operations, declared_totals=totals,
                     section_i_reference=reference, computed_totals={})
    report.data["extracted"] = extracted
    if title.get("tax_object", {}).get("value", "").lower() != "доходы" or tax_regime != "USN_INCOME":
        report.issue("UNSUPPORTED", "Only sole proprietor USN income is supported")
    identity_ok = title.get("inn", {}).get("value") == inn
    report.check("BORROWER_INN", "PASS" if identity_ok else "FAIL", [title["inn"]["source"]] if "inn" in title else [])
    if not identity_ok:
        report.issue("INN_MISMATCH", "Title INN differs from applicant INN")
    incomplete = False
    columns = None
    pending = []
    in_reference = False

    def assign_pending(quarter):
        nonlocal pending
        for op in pending:
            op["quarter"] = quarter
            if op.get("date") and (int(op["date"][5:7])-1)//3+1 != quarter:
                report.issue("QUARTER_MISMATCH", "Operation date differs from block quarter", list(op["sources"].values()))
        pending = []

    for s in sheets:
        # Other sections must not inherit section I column positions.
        has_section = any(re.search(r"^I\.\s*Доходы", norm(v)) for row in s.rows for v in row.values())
        has_header = any("Дата и номер первичного" in " ".join(map(norm,row.values())) and "Содержание операции" in " ".join(map(norm,row.values())) for row in s.rows)
        continuation = columns and any("продолжение" in norm(v).lower() and re.search(r"раздела?\s+i\b",norm(v),re.I) for row in s.rows for v in row.values())
        if not has_section and not has_header and not continuation:
            continue
        in_reference = False
        for r, row in enumerate(s.rows):
            text = " ".join(norm(v) for _, v in sorted(row.items()))
            low = text.lower()
            if re.match(r"^I\.\s*Доходы", text) or ("регистрация" in low and "сумма" in low) or ("продолжение" in low and re.search(r"раздела?\s+i\b",low)):
                continue
            if "справка к разделу" in low:
                in_reference = True
                continue
            if in_reference:
                code = next((norm(v) for v in row.values() if norm(v) in {"010", "020", "030", "040", "041"}), None)
                if code:
                    # Values occur after descriptive text, not the row code itself.
                    label_col = max((c for c,v in row.items() if re.search(r"[А-Яа-я]", str(v))), default=-1)
                    nums = [(c,v) for c,v in row.items() if c > label_col and norm(v) not in MISSING and norm(v) != code]
                    if nums:
                        c, v = nums[-1]
                        try:
                            reference[code] = {"amount": money(v), "source": s.source(r,c)}
                        except ValueError:
                            report.issue("INVALID_AMOUNT", "Invalid section I reference", [s.source(r,c)])
                continue
            if "дата и номер первичного" in low and "доходы" in low:
                columns = {}
                for c, v in row.items():
                    label = norm(v).lower()
                    if "п/п" in label: columns["number"] = c
                    if "дата и номер первичного" in label: columns["document"] = c
                    if "содержание операции" in label: columns["description"] = c
                    if "доходы" in label and "налогов" in label: columns["income"] = c
                if len(columns) != 4:
                    report.issue("HEADER_UNSUPPORTED", "Section I columns are incomplete", [s.source(r,c) for c in row])
                    columns = None
                continue
            if not columns:
                continue
            key = total_key(low)
            if key:
                c = columns["income"]
                try:
                    amount = money(row.get(c))
                    if amount is not None:
                        item = {"amount": amount, "source": s.source(r,c)}
                        if key in totals:
                            report.issue("DUPLICATE_TOTAL", "Repeated total " + key, [totals[key]["source"], item["source"]])
                        else:
                            totals[key] = item
                except ValueError:
                    report.issue("INVALID_AMOUNT", "Invalid declared total", [s.source(r,c)])
                if key.startswith("Q"):
                    assign_pending(int(key[1]))
                continue
            vals = {k: row.get(c, "") for k,c in columns.items()}
            if any(norm(v) in {"…", "..."} for v in vals.values()):
                incomplete = True
                report.issue("INCOMPLETE_DETAIL", "Ellipsis denotes omitted operations", [s.source(r,c) for c in columns.values()], "WARNING")
                continue
            if all(norm(v) == str(i) for v,i in zip(vals.values(), (1,2,3,4))):
                continue
            if not any(norm(v) for v in vals.values()):
                continue
            number = re.fullmatch(r"\d+(?:\.0)?", str(vals["number"]))
            if not number:
                report.issue("UNRECOGNIZED_OPERATION", "Nonempty section I row is not an operation", [s.source(r,c) for c in row])
                continue
            try:
                amount = money(vals["income"])
            except ValueError:
                amount = None
            raw_doc = str(vals["document"])
            match = re.search(r"(\d{2}\.\d{2}\.\d{4})", raw_doc)
            date = None
            if match:
                try: date = dt.datetime.strptime(match[1], "%d.%m.%Y").date().isoformat()
                except ValueError: pass
            doc_num = re.search(r"(?:№|No\.?|N)[ \t]*(\S+)", raw_doc, re.I)
            sources = {k: s.source(r,c) for k,c in columns.items()}
            op = dict(number=int(float(vals["number"])), primary_document_raw=raw_doc,
                      date=date, document_number=doc_num[1] if doc_num else None,
                      description=str(vals["description"]), income=amount, quarter=None, sources=sources)
            if not date or not doc_num or amount is None or not norm(vals["description"]):
                report.issue("OPERATION_INCOMPLETE", "Operation contains missing/unrecognized fields", list(sources.values()))
            if date and int(date[:4]) != title.get("year", {}).get("value"):
                report.issue("YEAR_MISMATCH", "Operation year differs from title", list(sources.values()))
            operations.append(op)
            pending.append(op)
    if pending:
        report.issue("QUARTER_UNKNOWN", "Operations lack a closing quarterly total", [o["sources"]["income"] for o in pending])
    extracted["detail_complete"] = not incomplete
    if incomplete:
        report.issue("DETAIL_REQUIRED_FOR_SCORING", "Incomplete operation detail blocks automatic income scoring")
    for q in range(1,5):
        amounts = [Decimal(o["income"]) for o in operations if o["quarter"] == q and o["income"] is not None]
        extracted["computed_totals"]["Q"+str(q)] = format(sum(amounts), ".2f") if amounts else None
    extracted["computed_totals"]["observed_income"] = format(sum(Decimal(o["income"]) for o in operations if o["income"] is not None), ".2f") if operations else None
    for key in ("Q1", "Q2", "Q3", "Q4", "H1", "M9", "YEAR"):
        if key not in totals:
            report.issue("TOTAL_MISSING", "Missing declared total: " + key)
    for key, n in (("H1",2), ("M9",3), ("YEAR",4)):
        deps = ["Q"+str(q) for q in range(1,n+1)]
        if key in totals and all(k in totals for k in deps):
            ok = sum(Decimal(totals[k]["amount"]) for k in deps) == Decimal(totals[key]["amount"])
            sources = [totals[k]["source"] for k in deps+[key]]
            report.check("DECLARED_"+key, "PASS" if ok else "FAIL", sources)
            if not ok: report.issue("TOTAL_MISMATCH", "Quarterly totals disagree with " + key, sources)
    for q in range(1,5):
        key = "Q"+str(q)
        observed = extracted["computed_totals"][key]
        status = "NOT_CHECKABLE"
        if not incomplete and observed is not None and key in totals:
            status = "PASS" if observed == totals[key]["amount"] else "FAIL"
        report.check("OPERATIONS_"+key, status)
        if status != "PASS" and not incomplete:
            report.issue("OPERATIONS_UNRECONCILED", "Cannot reconcile operations for " + key)
    if "YEAR" in totals and "year" in title:
        report.data["indicators"].append(dict(name="usn_taxable_income", period=str(title["year"]["value"]), period_kind="YEAR",
                                               amount=totals["YEAR"]["amount"], basis="DECLARED", sources=[totals["YEAR"]["source"]]))


INDICATORS = {"2110": "accounting_revenue", "2400": "net_profit", "1600": "total_assets",
              "1300": "equity", "1400": "long_term_liabilities", "1500": "short_term_liabilities"}


def pdf_lines(page):
    lines = []
    for word in sorted(page.extract_words(), key=lambda w: (round(w["top"]/3), w["x0"])):
        row = next((r for r in reversed(lines[-3:]) if abs(r[0]["top"]-word["top"]) < 3), None)
        if row is None: lines.append([word])
        else: row.append(word)
    return [sorted(row, key=lambda w:w["x0"]) for row in lines]


def parse_pdf(path, report, inn):
    import pdfplumber
    values, identities = [], []
    extracted = dict(kind="GIRBO", values=values, taxpayer_inn=None)
    report.data["extracted"] = extracted
    scale = None
    unit_source = None
    form = None
    forms = set()
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > 200: raise ValueError("PDF exceeds page limit")
        for pno,page in enumerate(pdf.pages,1):
            text = page.extract_text() or ""
            if not text.strip():
                report.issue("NO_TEXT_LAYER", "OCR is not supported", [{"page":pno}])
                continue
            for line in pdf_lines(page):
                line_text = " ".join(w["text"] for w in line)
                m = re.search(r"ИНН\s+(\d{10})\b", line_text)
                if m: identities.append((m[1], {"page":pno, "raw":line_text}))
            if re.search(r"Единица измерения\s+Тыс\.\s*руб", text, re.I): scale = Decimal(1000)
            elif re.search(r"Единица измерения\s+(?:руб\.?|рублей)\b", text, re.I): scale = Decimal(1)
            elif re.search(r"Единица измерения\s+Млн\.\s*руб", text, re.I): scale = Decimal(1_000_000)
            if "Единица измерения" in text:
                unit_source = {"page":pno,"raw":next((line for line in text.splitlines() if "Единица измерения" in line),"")}
            if pno == 1:
                name = re.search(r'(ООО\s+[«"].+?[»"])',text)
                extracted["taxpayer_name"] = name[1] if name else None
                if not name:
                    report.issue("LEGAL_FORM_UNCONFIRMED", "Cannot confirm OOO in report title", [{"page":1}])
            if "Бухгалтерский баланс" in text: form = "BALANCE"; forms.add(form)
            if "Отчет о финансовых результатах" in text or "Отчёт о финансовых результатах" in text:
                form = "FINANCIAL_RESULTS"; forms.add(form)
            if not form: continue
            lines = pdf_lines(page)
            # Column positions come from year headings, not fixed PDF coordinates.
            headers = []
            for line in lines:
                years = [w for w in line if re.fullmatch(r"20\d{2}", w["text"])]
                if len(years) >= 2:
                    headers = years
                    break
            if not headers:
                continue
            for line in lines:
                codes = [w for w in line if re.fullmatch(r"[12]\d{3}", w["text"]) and
                         not w["text"].startswith("20") and w["x0"] < headers[0]["x0"]]
                if not codes: continue
                if len(codes) != 1:
                    report.issue("AMBIGUOUS_ROW", "Multiple row codes", [{"page":pno}]); continue
                code = codes[0]
                words = [w for w in line if w["x0"] > code["x1"]]
                for idx,h in enumerate(headers):
                    center = (h["x0"]+h["x1"])/2
                    left = (center+(headers[idx-1]["x0"]+headers[idx-1]["x1"])/2)/2 if idx else code["x1"]
                    right = (center+(headers[idx+1]["x0"]+headers[idx+1]["x1"])/2)/2 if idx+1<len(headers) else page.width
                    cell = [w for w in words if left <= (w["x0"]+w["x1"])/2 < right]
                    raw = " ".join(w["text"] for w in cell)
                    source = {"page":pno, "row_code":code["text"], "raw":raw,
                              "bbox":[left, min(w["top"] for w in line), right, max(w["bottom"] for w in line)],
                              "period_header": {"page":pno,"raw":h["text"],"bbox":[h["x0"],h["top"],h["x1"],h["bottom"]]}}
                    try: amount = money(raw)
                    except ValueError:
                        report.issue("INVALID_AMOUNT", "Unrecognized PDF amount", [source]); amount = None
                    normalized = format(Decimal(amount)*scale, ".2f") if amount is not None and scale is not None else None
                    values.append(dict(form=form, code=code["text"], period=h["text"],
                                       period_kind="AS_OF" if form=="BALANCE" else "YEAR",
                                       amount=normalized, original_amount=amount,
                                       unit_multiplier=str(scale) if scale is not None else None, unit_source=unit_source, source=source))
    if scale is None: report.issue("UNIT_UNKNOWN", "Report unit is missing")
    extracted["taxpayer_inn"] = identities[0][0] if identities else None
    ok = bool(identities) and all(x[0]==inn for x in identities)
    report.check("BORROWER_INN", "PASS" if ok else "FAIL", [x[1] for x in identities])
    if not ok: report.issue("INN_MISMATCH", "Report INN missing or differs from applicant", [x[1] for x in identities])
    if forms != {"BALANCE", "FINANCIAL_RESULTS"}:
        report.issue("FORMS_MISSING", "Balance and financial results are required")
    for required_form in ("BALANCE", "FINANCIAL_RESULTS"):
        if not any(v["form"] == required_form for v in values):
            report.issue("FORM_TABLE_MISSING", "No readable table for " + required_form)
    index = {}
    for v in values:
        key = (v["form"], v["code"], v["period"])
        if key in index:
            report.issue("DUPLICATE_VALUE", "Repeated financial value", [index[key]["source"],v["source"]])
        index[key] = v
        if v["code"] in INDICATORS and v["amount"] is not None:
            report.data["indicators"].append(dict(name=INDICATORS[v["code"]], period=v["period"], period_kind=v["period_kind"],
                                                   amount=v["amount"], basis="DECLARED", sources=[v["source"]]))
    for year in sorted({v["period"] for v in values}):
        def check(form_name, name, codes, calc, affected):
            rows = [index.get((form_name,c,year)) for c in codes]
            if not all(r and r["amount"] is not None for r in rows):
                report.check(name+"_"+year, "NOT_CHECKABLE")
                if any(r is not None for r in rows):
                    report.issue("ARITHMETIC_NOT_CHECKABLE", name+" cannot be checked for "+year, [r["source"] for r in rows if r], indicators=affected, period=year)
                # Missing a required total blocks the related metric, not a different year.
                return
            nums = [Decimal(r["amount"]) for r in rows]
            valid = calc(nums)
            sources = [r["source"] for r in rows]
            report.check(name+"_"+year, "PASS" if valid else "FAIL", sources)
            if not valid:
                report.issue("ARITHMETIC_MISMATCH", name+" "+year, sources, indicators=affected, period=year)
        check("BALANCE", "BALANCE_EQUALITY", ["1600","1700"], lambda n:n[0]==n[1], ["total_assets","equity","long_term_liabilities","short_term_liabilities"])
        check("BALANCE", "ASSET_TOTAL", ["1100","1200","1600"], lambda n:n[0]+n[1]==n[2], ["total_assets"])
        check("BALANCE", "LIABILITY_TOTAL", ["1300","1400","1500","1700"], lambda n:sum(n[:3])==n[3], ["equity","long_term_liabilities","short_term_liabilities"])
        check("FINANCIAL_RESULTS", "GROSS_PROFIT", ["2110","2120","2100"], lambda n:n[0]+n[1]==n[2], ["accounting_revenue","net_profit"])
        # A discrepancy is flagged only when explicit tax/subtotal/net rows exist.
        # Optional adjustments are not silently assumed zero; an incomplete bridge is diagnostic.
        base = [index.get(("FINANCIAL_RESULTS",c,year)) for c in ("2300","2410","2400")]
        if all(r and r["amount"] is not None for r in base):
            before,tax,net = [Decimal(r["amount"]) for r in base]
            adjustments = [index.get(("FINANCIAL_RESULTS",c,year)) for c in ("2460",)]
            complete_bridge = all(r and r["amount"] is not None for r in adjustments)
            extra = sum(Decimal(r["amount"]) for r in adjustments) if complete_bridge else None
            status = "NOT_CHECKABLE" if extra is None else ("PASS" if before-abs(tax)+extra == net else "FAIL")
            sources = [r["source"] for r in base + adjustments if r]
            report.check("NET_PROFIT_BRIDGE_"+year, status, sources)
            if status != "PASS":
                report.issue("NET_PROFIT_UNRECONCILED", "Profit, tax and net profit require reconciliation for "+year,
                             sources, indicators=["net_profit"], period=year)
    if not values: report.issue("NO_FINANCIAL_VALUES", "No financial table values extracted")


def extract(request):
    report = Report()
    path = Path(request["path"])
    if path.stat().st_size > 20*1024*1024: raise ValueError("File exceeds 20 MB")
    report.data["document_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    typ = request.get("applicant_type")
    header = path.read_bytes()[:8]
    try:
        if typ == "SOLE_PROPRIETOR" and (header == bytes.fromhex("d0cf11e0a1b11ae1") or header.startswith(b"PK")):
            parse_kudir(path, report, request["inn"], request.get("tax_regime"))
        elif typ == "LEGAL_ENTITY" and header.startswith(b"%PDF"):
            if request.get("legal_form") != "OOO":
                report.issue("UNSUPPORTED", "Only OOO is supported")
            else:
                parse_pdf(path, report, request["inn"])
        else:
            report.issue("UNSUPPORTED", "Unsupported applicant/document combination")
    except Exception as exc:
        report.issue("PARSE_FAILED", str(exc))
    return report.finish()


if __name__ == "__main__":
    # Subprocess resource limits apply before opening untrusted documents.
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (768*1024*1024,768*1024*1024))
    resource.setrlimit(resource.RLIMIT_CPU, (30,30))
    try:
        result = extract(json.load(sys.stdin))
    except Exception as exc:
        r = Report(); r.issue("PARSE_FAILED", str(exc)); result = r.finish()
    json.dump(result, sys.stdout, ensure_ascii=False, allow_nan=False)
